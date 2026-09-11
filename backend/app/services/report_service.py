import io
import csv
from datetime import date, datetime
from typing import List, Optional, Dict, Any
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_, func
from sqlalchemy.orm import selectinload

from app.models.attendance import AttendanceRecord, AttendanceSession
from app.models.student import Student
from app.models.academic import Department, Course, Subject
from app.schemas.reports import (
    ReportFilterParams,
    ReportRecordItem,
    StudentAttendanceSummary,
    AttendanceReportResponse
)

# Document generation libraries
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


class ReportService:

    @staticmethod
    async def generate_report_data(
        db: AsyncSession,
        filters: ReportFilterParams
    ) -> AttendanceReportResponse:
        # 1. Base Query with relationships
        query = (
            select(AttendanceRecord)
            .join(AttendanceSession, AttendanceRecord.session_id == AttendanceSession.id)
            .join(Student, AttendanceRecord.student_id == Student.id)
            .options(
                selectinload(AttendanceRecord.student).selectinload(Student.department),
                selectinload(AttendanceRecord.student).selectinload(Student.course),
                selectinload(AttendanceRecord.session).selectinload(AttendanceSession.subject),
                selectinload(AttendanceRecord.session).selectinload(AttendanceSession.course)
            )
            .order_by(AttendanceRecord.marked_at.desc())
        )

        conditions = []
        if filters.start_date:
            conditions.append(AttendanceSession.session_date >= filters.start_date)
        if filters.end_date:
            conditions.append(AttendanceSession.session_date <= filters.end_date)
        if filters.subject_id:
            conditions.append(AttendanceSession.subject_id == filters.subject_id)
        if filters.course_id:
            conditions.append(AttendanceSession.course_id == filters.course_id)
        if filters.department_id:
            conditions.append(Student.department_id == filters.department_id)
        if filters.semester:
            conditions.append(AttendanceSession.semester == filters.semester)
        if filters.section:
            conditions.append(AttendanceSession.section == filters.section.upper())
        if filters.student_id:
            conditions.append(AttendanceRecord.student_id == filters.student_id)
        if filters.status:
            conditions.append(AttendanceRecord.status == filters.status.upper())

        if conditions:
            query = query.where(and_(*conditions))

        result = await db.execute(query)
        records = result.scalars().all()

        # 2. Map detailed records
        record_items: List[ReportRecordItem] = []
        for r in records:
            record_items.append(ReportRecordItem(
                id=r.id,
                session_code=r.session.session_code if r.session else "SESS",
                session_date=r.session.session_date if r.session else date.today(),
                subject_code=r.session.subject.code if (r.session and r.session.subject) else "N/A",
                subject_name=r.session.subject.name if (r.session and r.session.subject) else "Lecture",
                course_name=r.session.course.name if (r.session and r.session.course) else "N/A",
                semester=r.session.semester if r.session else 1,
                section=r.session.section if r.session else "A",
                student_id=r.student.id,
                roll_number=r.student.student_id,
                student_name=r.student.name,
                department_name=r.student.department.name if (r.student and r.student.department) else "N/A",
                marked_at=r.marked_at,
                status=r.status,
                verification_method=r.verification_method,
                confidence_score=r.confidence_score,
                remarks=r.remarks
            ))

        # 3. Query total sessions in scope for attendance percentage calculation
        sess_query = select(func.count(AttendanceSession.id))
        sess_conditions = []
        if filters.start_date:
            sess_conditions.append(AttendanceSession.session_date >= filters.start_date)
        if filters.end_date:
            sess_conditions.append(AttendanceSession.session_date <= filters.end_date)
        if filters.subject_id:
            sess_conditions.append(AttendanceSession.subject_id == filters.subject_id)
        if filters.course_id:
            sess_conditions.append(AttendanceSession.course_id == filters.course_id)
        if filters.semester:
            sess_conditions.append(AttendanceSession.semester == filters.semester)
        if filters.section:
            sess_conditions.append(AttendanceSession.section == filters.section.upper())

        if sess_conditions:
            sess_query = sess_query.where(and_(*sess_conditions))
        sess_res = await db.execute(sess_query)
        total_sessions = sess_res.scalar_one() or 0
        if total_sessions == 0:
            # Fallback to unique session IDs in records
            total_sessions = len(set(r.session_id for r in records))

        # 4. Group by student to generate per-student summaries
        student_records_map: Dict[int, List[ReportRecordItem]] = {}
        student_meta_map: Dict[int, Dict[str, Any]] = {}

        for item in record_items:
            if item.student_id not in student_records_map:
                student_records_map[item.student_id] = []
                student_meta_map[item.student_id] = {
                    "student_id": item.student_id,
                    "roll_number": item.roll_number,
                    "student_name": item.student_name,
                    "department_name": item.department_name,
                    "course_name": item.course_name,
                    "semester": item.semester,
                    "section": item.section
                }
            student_records_map[item.student_id].append(item)

        student_summaries: List[StudentAttendanceSummary] = []
        for sid, recs in student_records_map.items():
            meta = student_meta_map[sid]
            p_count = sum(1 for r in recs if r.status == "PRESENT")
            l_count = sum(1 for r in recs if r.status == "LATE")
            attended = p_count + l_count

            # Calculate rate
            effective_sessions = max(total_sessions, attended)
            pct = round((attended / effective_sessions * 100), 1) if effective_sessions > 0 else 100.0

            student_summaries.append(StudentAttendanceSummary(
                student_id=meta["student_id"],
                roll_number=meta["roll_number"],
                student_name=meta["student_name"],
                department_name=meta["department_name"],
                course_name=meta["course_name"],
                semester=meta["semester"],
                section=meta["section"],
                total_sessions=effective_sessions,
                present_count=p_count,
                late_count=l_count,
                attendance_percentage=pct,
                status_warning=(pct < 75.0)
            ))

        student_summaries.sort(key=lambda s: s.roll_number)

        avg_pct = (
            round(sum(s.attendance_percentage for s in student_summaries) / len(student_summaries), 1)
            if student_summaries else 100.0
        )

        return AttendanceReportResponse(
            total_records=len(record_items),
            total_sessions=total_sessions,
            unique_students=len(student_summaries),
            average_attendance_percentage=avg_pct,
            records=record_items,
            student_summaries=student_summaries
        )

    @staticmethod
    def export_pdf(report: AttendanceReportResponse, title: str = "OFFICIAL COLLEGE ATTENDANCE REPORT") -> bytes:
        """Generates an institutional quality PDF document using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        normal_style = styles["Normal"]

        title_style = ParagraphStyle(
            "CollegeTitle",
            parent=styles["Heading1"],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#0f172a"),
            alignment=1,  # Centered
            fontName="Helvetica-Bold"
        )
        subtitle_style = ParagraphStyle(
            "CollegeSubtitle",
            parent=styles["Normal"],
            fontSize=10,
            leading=13,
            textColor=colors.HexColor("#475569"),
            alignment=1,
            fontName="Helvetica"
        )
        meta_label_style = ParagraphStyle(
            "MetaLabel",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#334155"),
            fontName="Helvetica-Bold"
        )

        elements = []

        # 1. College Header
        elements.append(Paragraph("COLLEGE OF ENGINEERING & TECHNOLOGY", title_style))
        elements.append(Paragraph("Department of Computer Science & Academic Affairs", subtitle_style))
        elements.append(Spacer(1, 4))
        elements.append(Paragraph(title, ParagraphStyle(
            "DocTitle", parent=subtitle_style, fontSize=11, fontName="Helvetica-Bold", textColor=colors.HexColor("#312e81")
        )))
        elements.append(Spacer(1, 10))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#4338ca"), spaceBefore=2, spaceAfter=8))

        # 2. Metadata Summary Bar
        now_str = datetime.now().strftime("%d %b %Y, %I:%M %p")
        meta_data = [
            [
                Paragraph(f"<b>Generated:</b> {now_str}", normal_style),
                Paragraph(f"<b>Total Sessions:</b> {report.total_sessions}", normal_style),
                Paragraph(f"<b>Total Marked Records:</b> {report.total_records}", normal_style),
            ],
            [
                Paragraph(f"<b>Unique Students:</b> {report.unique_students}", normal_style),
                Paragraph(f"<b>Class Average Attendance:</b> {report.average_attendance_percentage}%", normal_style),
                Paragraph("<b>Biometric Verification:</b> YuNet + SFace", normal_style),
            ]
        ]
        meta_table = Table(meta_data, colWidths=[180, 180, 180])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ('INNERGRID', (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ]))
        elements.append(meta_table)
        elements.append(Spacer(1, 14))

        # 3. Student Attendance Summary Table
        elements.append(Paragraph("Student Attendance Roster & Presence Percentages", ParagraphStyle(
            "SectionH", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor("#0f172a")
        )))
        elements.append(Spacer(1, 6))

        table_rows = [
            ["Roll Number", "Student Name", "Dept", "Sem/Sec", "Attended", "Late", "Rate %", "Status"]
        ]

        for s in report.student_summaries:
            status_text = "WARNING" if s.status_warning else "REGULAR"
            table_rows.append([
                s.roll_number,
                s.student_name,
                s.department_name[:12],
                f"S{s.semester}-{s.section}",
                f"{s.present_count + s.late_count}/{s.total_sessions}",
                str(s.late_count),
                f"{s.attendance_percentage}%",
                status_text
            ])

        # If empty
        if len(table_rows) == 1:
            table_rows.append(["--", "No attendance data found", "--", "--", "--", "--", "--", "--"])

        t_summary = Table(
            table_rows,
            colWidths=[80, 130, 65, 55, 60, 45, 55, 50],
            repeatRows=1
        )
        t_summary.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e293b")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8.5),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('TOPPADDING', (0, 0), (-1, 0), 6),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (4, 0), (6, -1), 'CENTER'),
            ('ALIGN', (7, 0), (7, -1), 'CENTER'),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ('TOPPADDING', (0, 1), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 1), (-1, -1), 4),
        ]))
        elements.append(t_summary)
        elements.append(Spacer(1, 30))

        # 4. Official Signatures
        sig_data = [
            [
                Paragraph("__________________________<br/><b>Faculty In-Charge</b><br/>Subject Coordinator", normal_style),
                Paragraph("__________________________<br/><b>Verified By</b><br/>Academic Dean / HOD", normal_style),
                Paragraph("__________________________<br/><b>System Seal</b><br/>Antigravity Biometrics", normal_style)
            ]
        ]
        sig_table = Table(sig_data, colWidths=[180, 180, 180])
        sig_table.setStyle(TableStyle([
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        elements.append(sig_table)

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_excel(report: AttendanceReportResponse) -> bytes:
        """Generates a multi-sheet, beautifully formatted Excel spreadsheet (.xlsx)."""
        wb = Workbook()

        # Styles
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        border_thin = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")

        # --- Sheet 1: Student Attendance Summary ---
        ws1 = wb.active
        ws1.title = "Attendance Summary"
        ws1.views.sheetView[0].showGridLines = True

        headers1 = [
            "Roll Number", "Student Name", "Department", "Course",
            "Semester", "Section", "Total Sessions", "Present Count",
            "Late Count", "Attendance %", "Deficit Warning (<75%)"
        ]
        ws1.append(headers1)

        for col_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        warning_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        warning_font = Font(name="Arial", size=9, bold=True, color="991B1B")

        for row_idx, s in enumerate(report.student_summaries, start=2):
            ws1.append([
                s.roll_number,
                s.student_name,
                s.department_name,
                s.course_name,
                s.semester,
                s.section,
                s.total_sessions,
                s.present_count,
                s.late_count,
                f"{s.attendance_percentage}%",
                "YES" if s.status_warning else "NO"
            ])

            for col_idx in range(1, len(headers1) + 1):
                cell = ws1.cell(row=row_idx, column=col_idx)
                cell.border = border_thin
                cell.font = Font(name="Arial", size=9)
                if col_idx in [5, 6, 7, 8, 9, 10, 11]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

                if s.status_warning and col_idx in [10, 11]:
                    cell.fill = warning_fill
                    cell.font = warning_font

        # Auto-adjust column widths for Sheet 1
        for col in ws1.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # --- Sheet 2: Detailed Attendance Logs ---
        ws2 = wb.create_sheet(title="Detailed Attendance Logs")
        ws2.views.sheetView[0].showGridLines = True

        headers2 = [
            "Record ID", "Session Code", "Date", "Subject Code",
            "Subject Name", "Course", "Roll Number", "Student Name",
            "Marked At", "Status", "Method", "Confidence", "Remarks"
        ]
        ws2.append(headers2)

        for col_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        for row_idx, r in enumerate(report.records, start=2):
            ws2.append([
                r.id,
                r.session_code,
                r.session_date.strftime("%Y-%m-%d"),
                r.subject_code,
                r.subject_name,
                r.course_name,
                r.roll_number,
                r.student_name,
                r.marked_at.strftime("%Y-%m-%d %H:%M:%S"),
                r.status,
                r.verification_method,
                f"{r.confidence_score*100:.1f}%" if r.confidence_score else "N/A",
                r.remarks or ""
            ])

            for col_idx in range(1, len(headers2) + 1):
                cell = ws2.cell(row=row_idx, column=col_idx)
                cell.border = border_thin
                cell.font = Font(name="Arial", size=9)
                if col_idx in [1, 3, 4, 10, 11, 12]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        # Auto-adjust column widths for Sheet 2
        for col in ws2.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def export_csv(report: AttendanceReportResponse) -> str:
        """Generates standard RFC 4180 CSV with UTF-8 BOM for universal Excel/ERP import."""
        output = io.StringIO()
        # Add UTF-8 BOM so Excel opens CSV with proper encoding
        output.write('\ufeff')
        writer = csv.writer(output)

        writer.writerow([
            "Record ID", "Session Code", "Date", "Subject Code",
            "Subject Name", "Course", "Semester", "Section",
            "Roll Number", "Student Name", "Department",
            "Marked At", "Status", "Verification Method", "Confidence Score", "Remarks"
        ])

        for r in report.records:
            writer.writerow([
                r.id,
                r.session_code,
                r.session_date.strftime("%Y-%m-%d"),
                r.subject_code,
                r.subject_name,
                r.course_name,
                r.semester,
                r.section,
                r.roll_number,
                r.student_name,
                r.department_name,
                r.marked_at.strftime("%Y-%m-%d %H:%M:%S"),
                r.status,
                r.verification_method,
                f"{r.confidence_score:.3f}" if r.confidence_score else "",
                r.remarks or ""
            ])

        return output.getvalue()
